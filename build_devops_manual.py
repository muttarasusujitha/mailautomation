import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
DATASET_DIR = ROOT / "microservices/services/trainer-service/app/datasets_compact"
OUTPUT_DIR = ROOT / "outputs/training_manuals"

NAVY = RGBColor(11, 37, 69)
BLUE = RGBColor(46, 116, 181)
MUTED = RGBColor(89, 99, 110)


def set_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    props.append(shading)


def set_cell_width(cell, width):
    cell.width = Inches(width)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    set_font(p.add_run(text), 10.5)


def add_label_paragraph(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    set_font(p.add_run(label + ": "), 10.5, NAVY, True)
    set_font(p.add_run(text), 10.5)


def outcomes_for(day):
    tools = day["tools"]
    subtopics = day["subtopics"]
    primary = tools[0]
    secondary = tools[1] if len(tools) > 1 else primary
    return [
        f"Use {primary} to work through {subtopics[0]} and {subtopics[1] if len(subtopics) > 1 else subtopics[0]}.",
        f"Apply {secondary} during the lab to produce the stated day deliverable.",
        f"Record evidence for '{day.get('jira_focus') or 'the completed lab activity'}' in the Jira sprint board.",
    ]


def configure(doc, domain):
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 14, 7),
        ("Heading 2", 13, NAVY, 10, 5),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run(f"{domain} Mastery | 30-Day Training Manual"), 8.5, MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("Trainer reference - structured from the approved DevOps curriculum"), 8.5, MUTED)


def build_manual(dataset_path):
    data = json.loads(dataset_path.read_text(encoding="utf-8"))
    domain = str(data.get("name") or data.get("domain") or dataset_path.stem.replace("_30", "").replace("_", " ").title())
    programme_tools = list(data.get("tools") or [])
    if not programme_tools:
        for item in data.get("days") or []:
            for tool in item.get("tools") or []:
                if tool not in programme_tools:
                    programme_tools.append(tool)
    output = OUTPUT_DIR / f"{dataset_path.stem.replace('_30', '')}_30_day_training_manual.docx"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure(doc, domain)

    doc.add_paragraph().paragraph_format.space_after = Pt(85)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(title.add_run(f"{domain.upper()} MASTERY"), 28, NAVY, True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(subtitle.add_run("30-Day Training Manual"), 16, BLUE)
    doc.add_paragraph().paragraph_format.space_after = Pt(70)
    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(intro.add_run("Intermediate | Online | Trainer and learner reference"), 11, MUTED)
    doc.add_page_break()

    doc.add_heading("Programme Guide", level=1)
    add_label_paragraph(doc, "Purpose", f"A practical 30-day {domain} programme structured around the approved daily curriculum, guided tools practice, Jira evidence, and a final project or review.")
    doc.add_heading("Core Tooling", level=2)
    for tool in programme_tools or [domain]:
        add_bullet(doc, tool)
    doc.add_heading("Certification Roadmap", level=2)
    for certification in data.get("certifications") or [f"Relevant {domain} certification roadmap"]:
        add_bullet(doc, certification)
    doc.add_heading("Daily Delivery Pattern", level=2)
    for item in ["Review the daily topic and subtopics.", "Demonstrate the named tools in a guided environment.", "Complete the day-specific lab and retain evidence.", "Update the Jira item, status, or release evidence described for the day."]:
        add_bullet(doc, item)

    for day in data["days"]:
        doc.add_page_break()
        doc.add_heading(f"Day {day['day']} - {day['topic']}", level=1)
        add_label_paragraph(doc, "Tools", ", ".join(day["tools"]))
        doc.add_heading("Session Focus", level=2)
        for subtopic in day["subtopics"]:
            add_bullet(doc, subtopic)
        doc.add_heading("Learning Outcomes", level=2)
        for outcome in outcomes_for(day):
            add_bullet(doc, outcome)
        doc.add_heading("Hands-on Lab", level=2)
        add_label_paragraph(doc, "Activity", day.get("lab_task") or day.get("lab") or "Complete the guided hands-on activity for this module.")
        add_label_paragraph(doc, "Jira evidence", day.get("jira_focus") or "Record the completed lab evidence and update the related task.")

    doc.add_page_break()
    doc.add_heading("Official Learning References", level=1)
    for source in data.get("official_sources") or []:
        add_bullet(doc, source)
    add_label_paragraph(doc, "Alignment note", data.get("official_alignment_note") or "Use the named tools and vendor documentation for future syllabus audits.")
    doc.core_properties.title = f"{domain} Mastery - 30-Day Training Manual"
    doc.core_properties.subject = "Trainer and learner reference"
    doc.core_properties.author = "Trainer Service"
    doc.save(output)
    return output


def main():
    datasets = sorted(DATASET_DIR.glob("*_30.json"))
    if not datasets:
        raise RuntimeError(f"No compact 30-day datasets found in {DATASET_DIR}")
    for dataset in datasets:
        print(build_manual(dataset))


if __name__ == "__main__":
    main()
