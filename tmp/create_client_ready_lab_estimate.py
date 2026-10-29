from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


source = Path(r"C:\Users\sujit\Downloads\DevOps - Lab Cost Estimate (1).xlsx")
destination = Path(r"C:\Users\sujit\Desktop\mail\outputs\DevOps - Lab Cost Estimate - Client Ready.xlsx")

book = openpyxl.load_workbook(source)
detail = book['Lab Cost Estimate']
if 'Client Summary' in book.sheetnames:
    del book['Client Summary']
summary = book.create_sheet('Client Summary', 0)

navy = '17365D'
teal = '0F766E'
light_teal = 'E6FFFB'
slate = '475569'
white = 'FFFFFF'
thin = Side(style='thin', color='CBD5E1')

summary.merge_cells('A1:F1')
summary['A1'] = 'DevOps Lab Cost Estimate'
summary['A1'].font = Font(name='Aptos Display', bold=True, size=18, color=white)
summary['A1'].fill = PatternFill('solid', fgColor=teal)
summary['A1'].alignment = Alignment(horizontal='left', vertical='center')
summary.row_dimensions[1].height = 34

summary.merge_cells('A2:F2')
summary['A2'] = 'Client quote summary — verified public retail rates. Provider alternatives are shown separately and must not be added together.'
summary['A2'].font = Font(name='Aptos', italic=True, size=10, color=slate)
summary['A2'].alignment = Alignment(wrap_text=True, vertical='center')
summary.row_dimensions[2].height = 30

summary.merge_cells('A4:C4')
summary['A4'] = 'AWS | Mumbai'
summary['A4'].font = Font(bold=True, color=white, size=12)
summary['A4'].fill = PatternFill('solid', fgColor=navy)
summary['A4'].alignment = Alignment(vertical='center')
summary.row_dimensions[4].height = 24

labels = [('Training title', "='Lab Cost Estimate'!B25"),
          ('Participants', "='Lab Cost Estimate'!B27"),
          ('Training duration', "='Lab Cost Estimate'!B28"),
          ('Infrastructure cost', "='Lab Cost Estimate'!B29"),
          ('Support cost', "='Lab Cost Estimate'!B30"),
          ('Contingency', "='Lab Cost Estimate'!B31"),
          ('Tax / GST', "='Lab Cost Estimate'!B32")]
for row, (label, formula) in enumerate(labels, 5):
    summary.cell(row, 1, label)
    summary.cell(row, 2, formula)
    summary.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
    summary.cell(row, 1).font = Font(bold=True, color=slate)
    summary.cell(row, 2).alignment = Alignment(horizontal='right')

summary['A12'] = 'Final estimated cost'
summary['B12'] = "='Lab Cost Estimate'!B33"
summary.merge_cells('B12:C12')
for cell in (summary['A12'], summary['B12']):
    cell.font = Font(bold=True, size=13, color=white)
    cell.fill = PatternFill('solid', fgColor=teal)
    cell.alignment = Alignment(horizontal='right' if cell.column == 2 else 'left', vertical='center')
summary.row_dimensions[12].height = 27
summary['A13'] = 'Cost per participant'
summary['B13'] = "='Lab Cost Estimate'!B34"
summary.merge_cells('B13:C13')
summary['A13'].font = Font(bold=True, color=slate)
summary['B13'].alignment = Alignment(horizontal='right')
summary['A14'] = 'Quote validity'
summary['B14'] = "='Lab Cost Estimate'!B35"
summary.merge_cells('B14:C14')

summary.merge_cells('A16:F16')
summary['A16'] = 'Included documents'
summary['A16'].font = Font(bold=True, color=white, size=11)
summary['A16'].fill = PatternFill('solid', fgColor=navy)
for row, text in enumerate(('Trainer profile', 'Training ToC', 'Detailed resource cost, assumptions, and verified rate card'), 17):
    summary.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    summary.cell(row, 1, '• ' + text)
    summary.cell(row, 1).fill = PatternFill('solid', fgColor=light_teal)

for row in range(5, 15):
    for column in range(1, 4):
        summary.cell(row, column).border = Border(bottom=thin)
for column, width in {'A': 28, 'B': 20, 'C': 16, 'D': 4, 'E': 24, 'F': 24}.items():
    summary.column_dimensions[column].width = width
for row in range(5, 15):
    summary.row_dimensions[row].height = 21
for cell in ('B8', 'B9', 'B10', 'B11', 'B12', 'B13'):
    summary[cell].number_format = '₹#,##0.00'
summary['B6'].number_format = '#,##0'
summary.sheet_view.showGridLines = False
summary.freeze_panes = 'A4'
summary.sheet_properties.pageSetUpPr.fitToPage = True
summary.page_setup.fitToWidth = 1
summary.page_setup.fitToHeight = 1
summary.print_area = 'A1:F19'

# Make the detailed worksheet easier to scan without changing its formulas.
detail.sheet_view.showGridLines = False
detail.freeze_panes = 'A4'
detail.row_dimensions[1].height = 28
detail.row_dimensions[2].height = 28
detail['A1'].font = Font(bold=True, color=white, size=16)
detail['A1'].fill = PatternFill('solid', fgColor=teal)
detail['A1'].alignment = Alignment(vertical='center')
detail['A2'].alignment = Alignment(wrap_text=True, vertical='center')
for col, width in {'A': 34, 'B': 18, 'C': 16, 'D': 16, 'E': 18, 'F': 14, 'G': 38, 'H': 18, 'I': 14, 'J': 18, 'K': 12}.items():
    detail.column_dimensions[col].width = width
for row in range(1, detail.max_row + 1):
    if row not in (1, 2):
        detail.row_dimensions[row].height = max(detail.row_dimensions[row].height or 20, 20)

destination.parent.mkdir(parents=True, exist_ok=True)
book.save(destination)
print(destination)
