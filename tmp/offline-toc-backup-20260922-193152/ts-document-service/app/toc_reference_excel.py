"""Reference-inspired client TOCs rendered from one validated curriculum."""
import io
import math
import textwrap

from shared.toc_layouts import REFERENCE_LAYOUTS


def render_reference_toc(toc):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    layout = toc["excel_layout"]
    if layout not in REFERENCE_LAYOUTS:
        raise ValueError("Unknown TOC layout")
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

    def text(value):
        if isinstance(value, list):
            return "\n".join(text(item) for item in value)
        if isinstance(value, dict):
            return text(value.get("topic") or value.get("title") or "")
        return str(value or "")

    def sheet(name, headers, rows, widths):
        ws = workbook.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append(row)
        for column, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(column)].width = width
        for row in ws:
            lines = 1
            for cell in row:
                # Client text is data, including strings beginning with '='.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.font = Font(name="Calibri", size=11, bold=cell.row == 1,
                                 color="FFFFFF" if cell.row == 1 else "172B4D")
                cell.fill = PatternFill("solid", fgColor="24476A" if cell.row == 1 else
                                        "EFF4F8" if cell.row % 2 == 0 else "FFFFFF")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                width = max(8, int(widths[cell.column - 1]) - 3)
                lines = max(lines, sum(max(1, len(textwrap.wrap(part, width)))
                                       for part in str(cell.value or "").split("\n")))
            ws.row_dimensions[row[0].row].height = min(409, max(32, 15 * lines + 12))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.print_title_rows = "1:1"
        ws.print_area = ws.dimensions
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_view.showGridLines = False
        return ws

    original_days = toc.get("days") or []
    days = []
    for day in original_days:
        if day.get("modules"):
            for module in day["modules"]:
                days.append({**module, "day": day.get("day"), "date": day.get("date"),
                    "timing": day.get("timing"), "focus_area": module.get("title") or module.get("topic"),
                    "hours_per_day": None})
        else:
            days.append(day)
    def hours(day):
        minutes = day.get("minutes")
        try:
            value = float(minutes) / 60 if minutes is not None else day.get("hours_per_day", toc.get("hours_per_day"))
            value = float(value)
            return value if math.isfinite(value) and value > 0 else None
        except (TypeError, ValueError):
            return None

    durations = [hours(day) for day in days]
    brief = toc.get("requirement_brief") or {}
    overview = sheet("Program Overview", ["Programme", "Details"], [
        ["Title", toc.get("title") or toc.get("domain") or "Training Programme"],
        ["Overview", text(toc.get("overview"))],
        ["Audience", brief.get("audience") or toc.get("audience_level") or toc.get("level") or "To be confirmed"],
        ["Training days", len(original_days)],
        ["Planned training hours", sum(durations) if all(v is not None for v in durations) else "To be confirmed"],
        ["Delivery mode", toc.get("mode") or "To be confirmed"],
        ["Prerequisites", text(toc.get("prerequisites")) or "To be confirmed"],
        ["Learning outcomes", text(toc.get("learning_outcomes"))],
    ], [28, 105])

    plan_rows, outcomes, assessments = [], [], []
    for index, day in enumerate(days, 1):
        title = day.get("focus_area") or day.get("module") or day.get("title") or f"Module {index}"
        topics = text(day.get("subtopics") or (day.get("morning_session") or {}).get("topics"))
        lab = text(day.get("lab") or day.get("lab_task"))
        objectives = text(day.get("learning_objectives"))
        assessment = text(day.get("assessment")) or "To be confirmed"
        deliverable = text(day.get("deliverable") or day.get("lab_deliverable")) or "To be confirmed"
        duration = durations[index - 1]
        row = [day.get("day") or index, day.get("date") or "To be confirmed",
               day.get("timing") or toc.get("timing") or "To be confirmed", title,
               duration, topics, lab, objectives, deliverable, assessment, text(day.get("tools"))]
        plan_rows.append(row)
        outcomes.append([row[0], title, objectives])
        assessments.append([row[0], title, assessment, deliverable])

    if layout == "execution_plan":
        schedule = sheet("Day-wise Plan", ["Day", "Date", "Session timing", "Module", "Hours", "Topics",
            "Hands-on lab", "Learning outcomes", "Deliverable", "Success measure"],
            [row[:10] for row in plan_rows], [8, 17, 24, 36, 10, 48, 50, 48, 36, 40])
    elif layout == "technical_plan":
        schedule = sheet("Day-wise Training Plan", ["Day", "Date", "Session timing", "Module", "Hours", "Technical coverage", "Practical lab", "Learning outcomes", "Tools"],
            [row[:8] + [row[10]] for row in plan_rows], [8, 17, 24, 36, 10, 58, 52, 48, 26])
    elif layout == "skills_matrix":
        schedule = sheet("Skills Matrix", ["Day", "Date", "Session timing", "Module", "Hours", "Focus areas", "Practical application", "Participant outcomes", "Participant output"],
            [row[:9] for row in plan_rows], [8, 17, 24, 38, 10, 55, 50, 48, 38])
    else:
        schedule = sheet("Detailed Syllabus", ["Day", "Date", "Session timing", "Module", "Hours", "Detailed topics", "Lab", "Learning objectives", "Lab deliverable", "Assessment", "Tools"],
            plan_rows, [8, 17, 24, 38, 10, 58, 52, 48, 36, 40, 26])
    for row in range(2, len(days) + 2):
        schedule.cell(row, 5).number_format = "0.##"
    if days and all(value is not None for value in durations):
        overview.cell(6, 2, f"=SUM('{schedule.title}'!E2:E{len(days)+1})")
        overview.cell(6, 2).number_format = "0.##"
    sheet("Module Outcomes", ["Day", "Module", "Learning outcomes"], outcomes, [8, 42, 100])
    sheet("Assessment", ["Day", "Module", "Assessment / success measure", "Deliverable"], assessments, [8, 40, 85, 65])
    scenario_rows = []
    for row, day in zip(plan_rows, days):
        if day.get("scenarios"):
            for scenario in day["scenarios"]:
                scenario_rows.append([row[0], row[3], text(scenario.get("situation")),
                    text(scenario.get("activity")), text(scenario.get("outcome")), text(scenario.get("evidence"))])
        elif day.get("scenario"):
            scenario_rows.append([row[0], row[3], text(day["scenario"]), row[6], row[7], row[9]])
    if scenario_rows:
        sheet("Scenarios", ["Day", "Module", "Scenario", "Participant activity", "Expected outcome", "Evidence"],
              scenario_rows, [8, 38, 60, 60, 55, 45])
    readiness = []
    for item in toc.get("prerequisites") or []:
        readiness.append(["Prerequisite", text(item)])
    for item in toc.get("assumptions") or []:
        readiness.append(["Assumption to confirm", text(item)])
    for item in toc.get("clarification_questions") or []:
        readiness.append(["Clarification", text(item)])
    for item in brief.get("constraints") or []:
        readiness.append(["Client constraint", text(item)])
    if readiness:
        sheet("Readiness & Risks", ["Category", "Requirement / action"], readiness, [30, 110])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
