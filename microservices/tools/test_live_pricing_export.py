"""Run inside document-service (stdin) or with its app on PYTHONPATH."""
import asyncio
import io
from unittest.mock import AsyncMock, MagicMock, patch
import openpyxl
from fastapi import HTTPException
from app.routes.excel import export_toc_lab_cost_workbook
from shared.test_live_lab_pricing import LivePricingTests


async def main():
    fixtures = LivePricingTests()
    assumptions = fixtures.inputs()
    assumptions.update(cloud_region='Mumbai', participant_count=1,
                       hours_per_day=2, fx_rate=85, rate_snapshot_source='')
    payload = {'toc': {'domain': 'Test', 'days': [{'day': 1, 'topic': 'EC2 cloud lab'}]},
               'assumptions': assumptions}
    collection = MagicMock()
    collection.find_one = AsyncMock(return_value=None)
    collection.insert_one = AsyncMock()
    db = MagicMock()
    db.__getitem__.return_value = collection
    with patch('shared.live_lab_pricing._aws_catalog', return_value=fixtures.catalog()):
        response = await export_toc_lab_cost_workbook(payload, db)
    wb = openpyxl.load_workbook(io.BytesIO(response.body))
    assert wb['Assumptions']['B28'].value == response.headers['X-Lab-Cost-Quote-ID']
    assert wb['Live Price Check']['E2'].value == 0.12
    assert wb['Live Price Check'].max_row == 12
    saved = collection.insert_one.call_args.args[0]
    assert saved['rates']['VM Light']['rate'] == 0.12
    assert saved['toc'] == payload['toc']
    collection.find_one.return_value = saved
    with patch('shared.live_lab_pricing._aws_catalog', return_value=fixtures.catalog('0.15')):
        revised = await export_toc_lab_cost_workbook(payload, db)
    revised_wb = openpyxl.load_workbook(io.BytesIO(revised.body))
    assert revised_wb['Live Price Check']['I2'].value == 0.12
    assert round(revised_wb['Live Price Check']['J2'].value) == 25
    assert revised_wb['Live Price Check']['K2'].value == 'Yes'
    assert revised.headers['X-Lab-Cost-Quote-ID'] != response.headers['X-Lab-Cost-Quote-ID']
    collection.find_one.return_value = None
    collection.insert_one.reset_mock()
    payload['assumptions']['pricing_selections'] = {}
    try:
        await export_toc_lab_cost_workbook(payload, db)
        raise AssertionError('Missing mappings must not produce a workbook')
    except HTTPException as exc:
        assert exc.status_code == 422
    collection.insert_one.assert_not_called()
    print('Export integration passed: workbook rates, authoritative ID, audit record, missing-mapping block')


asyncio.run(main())
