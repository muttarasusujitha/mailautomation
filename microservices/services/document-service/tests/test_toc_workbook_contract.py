import io
import openpyxl
import pytest
from app.routes.excel import _toc_to_excel


@pytest.mark.parametrize('mode', ['template', 'ai'])
def test_export_preserves_scope_dates_and_template_format(mode):
    toc = {'generation_mode': mode, 'days': [{
        'day': 1, 'date': '2026-11-02', 'timing': '7:30 PM - 11:30 PM IST',
        'focus_area': 'DevOps CI/CD', 'subtopics': ['Git branching', 'Pipeline gates'],
        'lab': 'Build and validate a CI pipeline',
        'learning_objectives': ['Demonstrate successful deployment'],
    }]}
    wb = openpyxl.load_workbook(io.BytesIO(_toc_to_excel(toc)))
    sheet = wb['ToC Template'] if 'ToC Template' in wb.sheetnames else wb.active
    values = [cell.value for cell in sheet[2]]
    assert 'DevOps CI/CD' in values
    assert 'Git branching\nPipeline gates' in values
    assert '7:30 PM - 11:30 PM IST' in values
    assert sheet.cell(2, 2).value.year == 2026
    assert all(sheet.cell(2, col).alignment.wrap_text for col in range(1, 9))
    assert sheet.freeze_panes == 'A2'
