import io

import openpyxl
import pytest

from app.lab_sheet import combine_lab_estimates, _relocate_formula
from app.routes.excel import _lab_cost_to_excel


def test_relocation_preserves_absolute_and_range_references():
    assert _relocate_formula('=SUMPRODUCT(\'TOC Mapping\'!$C$4:$C$6)*\'Assumptions\'!B7+B2',
                             'Costs', {'TOC Mapping': 100, 'Assumptions': 200, 'Costs': 10}) == \
        '=SUMPRODUCT($C$104:$C$106)*B207+B12'


def test_external_references_are_rejected():
    with pytest.raises(ValueError):
        _relocate_formula("='[old.xlsx]Costs'!A1", 'Costs', {'Costs': 1})


def test_real_provider_templates_combine_with_local_formulas_only():
    inputs = []
    for provider, region in [('aws', 'ap-south-1'), ('azure', 'centralindia')]:
        content = _lab_cost_to_excel({'title': 'DevOps', 'days': [
            {'focus_area': 'Docker', 'subtopics': ['Build containers']},
        ]}, {'cloud_provider': provider, 'cloud_region': region,
             'participant_count': 1, 'hours_per_day': 3, 'fx_rate': 90})
        inputs.append((provider, content))
    result = openpyxl.load_workbook(io.BytesIO(combine_lab_estimates(inputs)))
    assert result.sheetnames == ['Lab Cost Estimate']
    sheet = result.active
    labels = [row[0].value for row in sheet]
    assert 'AWS — Client Estimate' in labels
    assert 'AZURE — Client Estimate' in labels
    assert 'AWS — Resource Cost Breakdown' in labels
    formulas = [cell.value for row in sheet for cell in row if cell.data_type == 'f']
    assert len(formulas) > 20
    assert all('!' not in formula and '#REF!' not in formula for formula in formulas)
    # Every provider retains the three-hour and single-learner assumptions.
    for provider in ('AWS', 'AZURE'):
        offset = labels.index(provider + ' — Assumptions') + 1
        assert sheet.cell(offset + 7, 2).value == 3
        assert sheet.cell(offset + 9, 2).value == 1


def test_single_provider_workbook_opens_on_client_estimate_and_uses_gb_months():
    content = _lab_cost_to_excel({'title': 'DevOps', 'days': [{'focus_area': 'Docker'}]}, {
        'cloud_provider': 'aws', 'cloud_region': 'ap-south-1',
        'participant_count': 1, 'hours_per_day': 3, 'fx_rate': 90,
    })
    workbook = openpyxl.load_workbook(io.BytesIO(content))
    assert workbook.active.title == 'Client Estimate'
    breakdown = workbook['Resource Cost Breakdown']
    assert '/30' in breakdown['B7'].value
    assert breakdown['H7'].value == '=B7*D7'
