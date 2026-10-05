import io

import openpyxl
import pytest

from app.routes.excel import _toc_to_excel
from shared.toc_layouts import REFERENCE_LAYOUTS


@pytest.mark.parametrize("layout", REFERENCE_LAYOUTS)
def test_reference_layout_preserves_validated_curriculum_and_numeric_hours(layout):
    toc = {"excel_layout": layout, "title": "Python for analysts", "hours_per_day": 4,
           "days": [{"day": 1, "date": "2026-11-02", "timing": "09:00-13:00", "minutes": 180,
                     "focus_area": "CSV validation", "subtopics": ["pandas", "Missing values"],
                     "lab": "Clean a sales CSV", "learning_objectives": ["Validate sales rows"],
                     "deliverable": "Cleaned CSV", "assessment": "Check all seeded errors",
                     "scenario": "A sales export has missing values", "tools": ["Python"]}]}
    wb = openpyxl.load_workbook(io.BytesIO(_toc_to_excel(toc)))
    plan = wb.worksheets[1]
    assert plan.cell(2, 5).value == 3
    assert plan.cell(2, 6).value == "pandas\nMissing values"
    assert plan.cell(2, 7).value == "Clean a sales CSV"
    assert plan.cell(2, 8).value == "Validate sales rows"
    assert plan.cell(2, 2).value == "2026-11-02"
    assert wb["Program Overview"]["B5"].value == 1
    assert wb["Program Overview"]["B6"].value == f"=SUM('{plan.title}'!E2:E2)"
    assert wb["Assessment"]["D2"].value == "Cleaned CSV"
    assert wb["Scenarios"]["C2"].value == "A sales export has missing values"
    assert all(sheet.freeze_panes == "A2" for sheet in wb)


def test_unknown_hours_are_not_invented_and_client_text_is_not_a_formula():
    wb = openpyxl.load_workbook(io.BytesIO(_toc_to_excel({"excel_layout": "skills_matrix", "days": [
        {"focus_area": "=HYPERLINK(\"https://example.com\")", "subtopics": ["SQL"]}]})))
    assert wb["Skills Matrix"]["D2"].data_type == "s"
    assert wb["Skills Matrix"]["E2"].value is None
    assert wb["Program Overview"]["B6"].value == "To be confirmed"
    assert wb["Training Delivery"]["D2"].value == "To be confirmed"
    assert wb["Training Delivery"]["E2"].value is None


@pytest.mark.parametrize("layout", REFERENCE_LAYOUTS)
def test_client_objectives_acceptance_and_preparation_survive_export(layout):
    toc = {"excel_layout": layout, "days": [{"day": 1, "modules": [{
        "title": "Sales reconciliation", "minutes": 90, "deliverable": "Reconciled sales report",
        "acceptance_checks": [{"input_or_condition": "Duplicated order ID", "expected_result": "Rejected once", "evidence": "Rejection report"}],
        "prerequisites": ["Python functions"], "preparation_requirements": ["Ten-row CSV with known errors"]}]}],
        "outcome_coverage": [{"client_outcome": "Reconcile finance totals", "modules": [{"day": 1,
            "module": "Sales reconciliation", "deliverable": "Reconciled sales report"}]}]}
    wb = openpyxl.load_workbook(io.BytesIO(_toc_to_excel(toc)))
    assert wb["Client Objectives"]["A2"].value == "Reconcile finance totals"
    assert wb["Acceptance Checks"]["D2"].value == "Rejected once"
    readiness = [r[1] for r in wb["Readiness & Risks"].iter_rows(min_row=2, values_only=True)]
    assert "Ten-row CSV with known errors" in readiness
    assert "Python functions" in readiness


@pytest.mark.parametrize("layout", REFERENCE_LAYOUTS)
def test_delivery_sequence_survives_every_layout(layout):
    toc = {"excel_layout": layout, "days": [{"day": 1, "modules": [{
        "title": "CSV validation", "minutes": 60, "timing": "10:00-11:00",
        "deliverable": "Cleaned CSV", "assessment": "Detect all seeded errors",
        "delivery_steps": [
            {"method": "Demonstration", "minutes": 15, "activity": "Inspect invalid rows", "evidence": "Error inventory"},
            {"method": "Practice and assessment", "minutes": 45, "activity": "Implement and test validation", "evidence": "Seeded errors detected"}
        ]}]}]}
    wb = openpyxl.load_workbook(io.BytesIO(_toc_to_excel(toc)))
    delivery = wb["Training Delivery"]
    assert delivery.max_row == 3
    assert delivery["E2"].value + delivery["E3"].value == 60
    assert delivery["F3"].value == "Implement and test validation"
    assert delivery["G3"].value == "Seeded errors detected"
    assert wb.worksheets[1]["C2"].value == "10:00-11:00"
    assert wb["Assessment"]["D2"].value == "Cleaned CSV"


@pytest.mark.parametrize("layout", REFERENCE_LAYOUTS)
def test_multiple_modules_per_day_preserve_total_and_outcomes(layout):
    toc = {"excel_layout": layout, "hours_per_day": 6, "days": [{"day": 1, "date": "2026-11-02",
        "modules": [
            {"title": "Requirements", "minutes": 90, "learning_objectives": ["Clarify scope"], "subtopics": ["Scope"]},
            {"title": "Implementation", "minutes": 120, "learning_objectives": ["Build a prototype"], "subtopics": ["Build"]},
        ]}]}
    wb = openpyxl.load_workbook(io.BytesIO(_toc_to_excel(toc)))
    schedule = wb.worksheets[1]
    assert schedule.max_row == 3
    assert schedule["E2"].value == 1.5
    assert schedule["E3"].value == 2
    assert schedule["H3"].value == "Build a prototype"
    assert wb["Program Overview"]["B5"].value == 1
    assert wb["Program Overview"]["B6"].value.endswith("!E2:E3)")


def test_reference_scenarios_keep_individual_evidence():
    toc = {"excel_layout": "execution_plan", "days": [{"day": 1, "modules": [{"title": "API Testing", "scenarios": [
        {"situation": "Missing token", "activity": "Call API", "outcome": "Access denied", "evidence": "401 response"},
        {"situation": "Invalid schema", "activity": "Validate payload", "outcome": "Reject payload", "evidence": "Schema report"},
    ]}]}]}
    wb = openpyxl.load_workbook(io.BytesIO(_toc_to_excel(toc)))
    assert wb["Scenarios"].max_row == 3
    assert wb["Scenarios"]["F3"].value == "Schema report"


def test_long_lab_is_visible_in_continuation_rows_without_duplicating_hours():
    long_lab = "\n".join(f"Lab step {index}: inspect input and capture results." for index in range(100))
    toc = {"excel_layout": "execution_plan", "days": [{"day": 1, "minutes": 240, "focus_area": "Long lab", "lab": long_lab}]}
    wb = openpyxl.load_workbook(io.BytesIO(_toc_to_excel(toc)))
    plan = wb["Day-wise Plan"]
    assert plan.max_row > 2
    assert sum(cell.value or 0 for cell in list(plan.columns)[4][1:]) == 4
    assert "Lab step 99" in "\n".join(str(row[6].value or "") for row in plan.iter_rows(min_row=2))
    assert wb["Program Overview"]["B6"].value.endswith(f"E{plan.max_row})")
