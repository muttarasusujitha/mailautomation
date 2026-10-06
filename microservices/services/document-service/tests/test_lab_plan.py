"""Regression coverage for the concise client view and combined exports."""
import io
import shutil

import openpyxl
import pytest

from app.lab_recalculation import recalculate_lab_workbook
from app.lab_sheet import combine_lab_estimates
from app.routes.excel import _lab_cost_to_excel


def sample_estimate(provider='aws', verified=True, **changes):
    toc_days = changes.pop('toc_days', None)
    # Synthetic rates/provenance for deterministic layout tests, never a quote.
    assumptions = {
        'cloud_provider': provider,
        'cloud_region': 'ap-south-1' if provider == 'aws' else 'centralindia',
        'participant_count': 10, 'hours_per_day': 3, 'fx_rate': 90,
        'vm_profile_rates': {'Light': .05, 'Heavy': .20},
        'vm_profile_sources': {'Light': 'https://example.com/test', 'Heavy': 'https://example.com/test'},
        'lab_day_mapping': [
            {'vm_qty': 10, 'vm_profile': 'Light', 'k8s_control_plane': 0,
             'k8s_worker_nodes': 0, 'managed_db': 0, 'object_storage_gb': 0, 'active_days': 1},
            {'vm_qty': 10, 'vm_profile': 'Heavy', 'k8s_control_plane': 1,
             'k8s_worker_nodes': 2, 'managed_db': 1, 'object_storage_gb': 12, 'active_days': 2},
            {'vm_qty': 0, 'vm_profile': 'None', 'k8s_control_plane': 0,
             'k8s_worker_nodes': 0, 'managed_db': 0, 'object_storage_gb': 0, 'active_days': 1},
        ],
    }
    if verified:
        assumptions.update(pricing_status='provider_api_verified_public_retail',
            rate_snapshot_id='synthetic-test-fixture', rate_checked_at='2026-09-21T00:00:00+00:00',
            rate_snapshot_source='https://example.com/test-only',
            fx_rate_source='https://example.com/test-only-fx', fx_rate_date='2026-09-21')
    assumptions.update(changes)
    return _lab_cost_to_excel({'title': 'TEST FIXTURE - NOT A CLIENT QUOTE', 'days': toc_days or [
        {'day': 1, 'focus_area': 'Linux', 'lab': 'Practice shell commands and file permissions'},
        {'day': '2-3', 'duration_days': 2, 'focus_area': 'Kubernetes', 'lab_task': 'Deploy and troubleshoot a sample service'},
        {'day': 4, 'focus_area': 'Local Docker', 'lab': '=1+1'},
    ]}, assumptions)


def test_module_lab_inputs_keep_teaching_and_access_hours_separate():
    days = [{'day': 1, 'modules': [
        {'title': 'Local practice', 'minutes': 60, 'lab': 'Docker Desktop on local machine'},
        {'title': 'Cloud practice', 'minutes': 90, 'lab': 'Build Linux server on AWS EC2'},
    ]}]
    from shared.lab_planning import default_cloud_mapping
    mapping = default_cloud_mapping({'days': days}, {'participant_count': 10})
    book = openpyxl.load_workbook(io.BytesIO(sample_estimate(toc_days=days, lab_day_mapping=mapping)))
    assert book['TOC Mapping']['C4'].value == 10
    assert book['Assumptions']['B7'].value == 3
    sheet = book['Module Lab Inputs']
    assert sheet['C2'].value == 60
    assert sheet['C3'].value == 90
    assert sheet['E2'].value == 'Local'
    assert sheet['F3'].value == "='Assumptions'!$B$7"
    assert 'Cloud practice' in book['Lab Plan']['C7'].value


def test_plan_preserves_toc_order_and_literal_lab_text():
    book = openpyxl.load_workbook(io.BytesIO(sample_estimate()))
    sheet = book.active
    assert sheet.title == 'Lab Plan'
    assert [sheet.cell(row, 1).value for row in (7, 8, 9)] == ['1', '2-3', '4']
    assert sheet['C8'].value == 'Deploy and troubleshoot a sample service'
    assert sheet['C9'].value == '=1+1'
    assert sheet['C9'].data_type == 's'
    assert 'F8' in sheet['G8'].value
    assert sheet['F8'].value == "='TOC Mapping'!I5*'Assumptions'!$B$7"


def test_unverified_rates_are_not_shown_as_numeric_costs():
    book = openpyxl.load_workbook(io.BytesIO(sample_estimate(verified=False)))
    sheet = book.active
    assert 'Incomplete' in sheet['B4'].value
    assert all(sheet.cell(row, 7).value == 'Unverified' for row in (7, 8, 9, 10, 11, 14, 19, 20))
    assert 'excluded' in sheet['B21'].value


def test_combined_view_shows_plans_and_keeps_literal_text():
    combined = combine_lab_estimates([(provider, sample_estimate(provider)) for provider in ('aws', 'azure')])
    sheet = openpyxl.load_workbook(io.BytesIO(combined)).active
    plans = [cell.row for cell in sheet['A'] if isinstance(cell.value, str) and cell.value.endswith(' — Lab Plan')]
    assert len(plans) == 2
    for offset in plans:
        assert not sheet.row_dimensions[offset + 7].hidden
        assert sheet.cell(offset + 9, 3).value == '=1+1'
        assert sheet.cell(offset + 9, 3).data_type == 's'
    quotes = [cell.row for cell in sheet['A'] if isinstance(cell.value, str) and cell.value.endswith(' — Client Estimate')]
    assert all(sheet.row_dimensions[row].hidden for row in quotes)


@pytest.mark.skipif(not (shutil.which('libreoffice') or shutil.which('soffice')), reason='Requires spreadsheet calculator')
def test_daily_cloud_costs_reconcile_after_recalculation():
    for provider in ('aws', 'azure'):
        raw = sample_estimate(provider)
        calculated = recalculate_lab_workbook(raw)
        book = openpyxl.load_workbook(io.BytesIO(calculated), data_only=True)
        plan = book['Lab Plan']
        assert plan['F8'].value == 6
        assert plan['G9'].value == 0  # Local day has no mapped cloud resources.
        assert plan['G10'].value + plan['G11'].value == pytest.approx(book['Client Estimate']['B8'].value)
        assert plan['G19'].value == pytest.approx(book['Client Estimate']['B12'].value)
        assert plan['G20'].value == pytest.approx(plan['G19'].value / 10)
        combined = recalculate_lab_workbook(combine_lab_estimates([(provider, raw)]))
        sheet = openpyxl.load_workbook(io.BytesIO(combined), data_only=True).active
        total = next(row[6].value for row in sheet if row[0].value == 'Priced lab subtotal - incomplete estimate')
        assert total == pytest.approx(plan['G19'].value)


def test_verified_cloud_rates_do_not_imply_complete_estimate():
    book = openpyxl.load_workbook(io.BytesIO(sample_estimate()))
    plan = book['Lab Plan']
    assert 'Incomplete estimate' in plan['B4'].value
    assert 'Verified cloud rates' in plan['B4'].value
    assert 'required licenses' in plan['B4'].value
    assert plan['A19'].value == 'Priced lab subtotal - incomplete estimate'
    assert plan['G19'].data_type == 'f'


def test_additional_live_meters_are_visible_and_included_in_workbook_total():
    raw = sample_estimate(
        pricing_selections={'Lambda invocations': {'unit_key': 'request'}},
        additional_resource_usage={'Lambda invocations': 1000},
        rate_card_overrides={'Lambda invocations': {
            'rate': 0.0000002, 'unit': 'Requests', 'sku': 'lambda-req',
            'source': 'https://pricing.example/lambda', 'verified_date': '2026-09-28',
        }},
    )
    book = openpyxl.load_workbook(io.BytesIO(raw))
    rate_card = book['Rate Card']
    rate_row = next(row for row in range(4, rate_card.max_row + 1)
                    if rate_card.cell(row, 3).value == 'Lambda invocations')
    breakdown = book['Resource Cost Breakdown']
    meter_row = next(row for row in range(4, breakdown.max_row + 1)
                     if breakdown.cell(row, 1).value == 'Lambda invocations')
    assert rate_card.cell(rate_row, 5).value == pytest.approx(0.0000002)
    assert rate_card.cell(rate_row, 7).value == 'https://pricing.example/lambda'
    assert breakdown.cell(meter_row, 2).value == 1000
    assert breakdown.cell(meter_row, 8).value == f'=B{meter_row}*D{meter_row}'
    assert breakdown[f'H{breakdown.max_row}'].value.startswith('=SUM(H4:H')


@pytest.mark.parametrize('settings,expected', [
    ({}, '=SUM(B8:B11)*(1+0.3)'),
    ({'clahan_margin_percent': 30}, '=SUM(B8:B11)*(1+0.3)'),
    ({'clahan_margin_percent': 15}, '=SUM(B8:B11)*(1+0.15)'),
    ({'clahan_margin_percent': 0}, '=SUM(B8:B11)*(1+0)'),
    ({'client_lab_markup_percent': 20}, '=SUM(B8:B11)*(1+0.2)'),
    ({'clahan_margin_percent': 10, 'client_lab_markup_percent': 20}, '=SUM(B8:B11)*(1+0.1)'),
])
def test_configured_clahan_markup_reaches_final_quote(settings, expected):
    book = openpyxl.load_workbook(io.BytesIO(sample_estimate(**settings)))
    assert book['Client Estimate']['B12'].value == expected


def test_confirmed_cost_scope_and_additional_unpriced_services():
    confirmed = dict(local_costs_status='covered', license_costs_status='not_required')
    complete = openpyxl.load_workbook(io.BytesIO(sample_estimate(**confirmed))).active
    assert complete['A19'].value == 'Total estimated lab cost (confirmed scope)'
    incomplete = openpyxl.load_workbook(io.BytesIO(sample_estimate(
        **confirmed, required_unpriced_costs=['GPU service']))).active
    assert 'GPU service' in incomplete['B4'].value
    assert 'incomplete' in incomplete['A19'].value


def test_vm_details_use_provider_metadata_and_flag_unknown_hardware():
    from app.lab_plan import _vm_description
    assert 'SKU unconfirmed' in _vm_description({}, 'Light')
    rates = {'VM Light': {'sku': 'synthetic', 'specifications': {
        'instance_type': 'test.small', 'vcpu': '2', 'memory': '4 GiB', 'os': 'Linux'}}}
    plan = openpyxl.load_workbook(io.BytesIO(sample_estimate(rate_card_overrides=rates))).active
    assert 'test.small / 2 vCPU / 4 GiB RAM / Linux' in plan['E7'].value
    assert 'vCPU unconfirmed' in plan['E7'].value  # Heavy remains unconfirmed.


def test_day_setup_labels_preserve_overrides_and_flag_local_cloud_conflicts():
    plan = openpyxl.load_workbook(io.BytesIO(sample_estimate(
        lab_day_setups=['individual', 'shared', 'local']))).active
    assert 'Setup: ' in plan['B7'].value and 'Individual' in plan['B7'].value
    assert 'Shared' in plan['B8'].value
    assert 'local setup with cloud resources; review' in plan['B9'].value
    inferred = openpyxl.load_workbook(io.BytesIO(sample_estimate())).active
    assert 'assumed' in inferred['B7'].value
