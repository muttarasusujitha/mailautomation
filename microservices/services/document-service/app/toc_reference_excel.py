"""Reference-inspired client TOCs rendered from one validated curriculum."""
import io
import math
import textwrap

from shared.toc_layouts import REFERENCE_LAYOUTS


# Ten natural shades. Each layout keeps the same curriculum and changes the stylesheet.
THEMES = {
    "execution_plan": {"header_bg": "FFF4EFE6", "header_fg": "FF3E3832", "band": "FFF7F1E8", "paper": "FFFFFCF8", "ink": "FF3E3832", "line": "FFE6D9C8", "rule": "FF8C7B6B", "label_bg": "FFF3E8D8", "label_fg": "FF3E3832", "tab": "8C7B6B", "style": "open", "header_align": "left", "header_size": 12, "zebra": True, "label_italic": True},
    "technical_plan": {"header_bg": "FF2F4A3C", "header_fg": "FFF7FBF7", "band": "FFE7F0E8", "paper": "FFF7FBF7", "ink": "FF24352C", "line": "FFD5E3D8", "rule": "FF7D9A84", "label_bg": "FFE7F0E8", "label_fg": "FF2F4A3C", "tab": "2F4A3C", "style": "grid", "header_align": "center", "header_size": 10, "zebra": True, "label_italic": False},
    "skills_matrix": {"header_bg": "FF8C4A3A", "header_fg": "FFFFF8F5", "band": "FFF8EBE3", "paper": "FFFBF6F2", "ink": "FF4A2E24", "line": "FFF0D9CE", "rule": "FFC4785A", "label_bg": "FFF8EBE3", "label_fg": "FF8C4A3A", "tab": "8C4A3A", "style": "accent", "header_align": "center", "header_size": 11, "zebra": True, "label_italic": False},
    "detailed_syllabus": {"header_bg": "FF1F4E5F", "header_fg": "FFF4FBFA", "band": "FFE5F2F2", "paper": "FFF7FBFA", "ink": "FF1A3338", "line": "FFD3E6E6", "rule": "FF6BA8A8", "label_bg": "FFE5F2F2", "label_fg": "FF1F4E5F", "tab": "1F4E5F", "style": "grid", "header_align": "center", "header_size": 12, "zebra": True, "label_italic": False},
    "stone_plan": {"header_bg": "FFE7E4DF", "header_fg": "FF3A403E", "band": "FFF3F1EE", "paper": "FFFAF9F7", "ink": "FF3A403E", "line": "FFE0DCD6", "rule": "FFA39E96", "label_bg": "FFE7E4DF", "label_fg": "FF3A403E", "tab": "6E726E", "style": "open", "header_align": "left", "header_size": 10, "zebra": True, "label_italic": False},
    "moss_plan": {"header_bg": "FF3D4F2F", "header_fg": "FFF8FBF3", "band": "FFE8F0DC", "paper": "FFF8FBF3", "ink": "FF2C3822", "line": "FFD7E4C8", "rule": "FF8FA36A", "label_bg": "FFE8F0DC", "label_fg": "FF3D4F2F", "tab": "3D4F2F", "style": "accent", "header_align": "left", "header_size": 11, "zebra": True, "label_italic": True},
    "sand_plan": {"header_bg": "FF7A6244", "header_fg": "FFFFFBF4", "band": "FFF8F1E3", "paper": "FFFDFBF6", "ink": "FF4A3B28", "line": "FFEFE2CC", "rule": "FFD4B483", "label_bg": "FFF8F1E3", "label_fg": "FF7A6244", "tab": "7A6244", "style": "ledger", "header_align": "center", "header_size": 10, "zebra": False, "label_italic": False},
    "bark_plan": {"header_bg": "FF4A3428", "header_fg": "FFF8F1EA", "band": "FFF3E6D8", "paper": "FFFBF7F2", "ink": "FF3A291F", "line": "FFE6D3C2", "rule": "FFA67C52", "label_bg": "FFF3E6D8", "label_fg": "FF4A3428", "tab": "4A3428", "style": "grid", "header_align": "left", "header_size": 12, "zebra": True, "label_italic": False},
    "mist_plan": {"header_bg": "FFE4EAF0", "header_fg": "FF2C3E50", "band": "FFEEF3F6", "paper": "FFF8FAFB", "ink": "FF2C3E50", "line": "FFD5DEE6", "rule": "FF8AA0B4", "label_bg": "FFE4EAF0", "label_fg": "FF2C3E50", "tab": "5C7388", "style": "open", "header_align": "center", "header_size": 11, "zebra": True, "label_italic": True},
    "olive_plan": {"header_bg": "FF556B2F", "header_fg": "FFF8FBEA", "band": "FFF0F3E4", "paper": "FFFBFCF6", "ink": "FF333E1E", "line": "FFE0E6CC", "rule": "FFA3B56A", "label_bg": "FFF0F3E4", "label_fg": "FF556B2F", "tab": "556B2F", "style": "ledger", "header_align": "left", "header_size": 11, "zebra": False, "label_italic": False},
}


def render_reference_toc(toc):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins

    layout = toc["excel_layout"]
    if layout not in REFERENCE_LAYOUTS or layout not in THEMES:
        raise ValueError("Unknown TOC layout")
    theme = THEMES[layout]
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    programme_title = str(toc.get("title") or toc.get("domain") or "Training Programme")
    if toc.get("draft_export"):
        programme_title = "DRAFT  ·  " + programme_title

    def text(value):
        if isinstance(value, list):
            return "\n".join(text(item) for item in value)
        if isinstance(value, dict):
            return text(value.get("topic") or value.get("title") or "")
        return str(value or "")

    def sheet(name, headers, rows, widths):
        ws = workbook.create_sheet(name)
        ws.sheet_properties.tabColor = theme["tab"]
        ws.append(headers)
        for row in rows:
            # Excel caps row height at 409 points. Spread long prose across
            # continuation rows instead of hiding the end of a detailed lab.
            chunks = []
            for index, value in enumerate(row):
                if not isinstance(value, str):
                    chunks.append([value])
                    continue
                wrapped = [line for paragraph in value.split("\n")
                           for line in (textwrap.wrap(paragraph, max(8, int(widths[index]) - 3)) or [""])]
                chunks.append(["\n".join(wrapped[start:start+22]) for start in range(0, len(wrapped), 22)] or [""]
                              if len(wrapped) > 22 else [value])
            for index in range(max(len(parts) for parts in chunks)):
                ws.append([parts[index] if index < len(parts) else None for parts in chunks])
        for column, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(column)].width = width
        hairline = Side(style="thin", color=theme["line"])
        rule = Side(style="medium", color=theme["rule"])
        boxed = theme["style"] in {"grid", "accent"}
        centered_headers = {"Day", "Hours", "Step", "Minutes"}
        centered_columns = {index for index, header in enumerate(headers, 1) if header in centered_headers}
        overview = name == "Program Overview"
        band = 0

        def edges(fill, bottom, left=None):
            edge = Side(style="thin", color=fill)
            return Border(
                left=left or (hairline if boxed else edge),
                right=hairline if boxed else edge,
                top=hairline if boxed else edge,
                bottom=bottom,
            )

        for row in ws:
            lines = 1
            starts_record = row[0].value not in (None, "")
            if row[0].row > 1 and starts_record:
                band += 1
            stripe = theme["paper"] if overview or not theme["zebra"] else (theme["band"] if band % 2 else theme["paper"])
            for cell in row:
                # Client text is data, including strings beginning with '='.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                if cell.row == 1:
                    cell.font = Font(name="Calibri", size=theme["header_size"], bold=True, color=theme["header_fg"])
                    cell.fill = PatternFill("solid", fgColor=theme["header_bg"])
                    cell.alignment = Alignment(horizontal=theme["header_align"], vertical="center", wrap_text=True, indent=1 if theme["header_align"] == "left" else 0)
                    cell.border = edges(theme["header_bg"], rule)
                elif overview and cell.column == 1:
                    cell.font = Font(name="Calibri", size=11, bold=True, italic=theme["label_italic"], color=theme["label_fg"])
                    cell.fill = PatternFill("solid", fgColor=theme["label_bg"])
                    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
                    cell.border = edges(theme["label_bg"], hairline, rule if theme["style"] == "accent" else None)
                else:
                    cell.font = Font(name="Calibri", size=11, color=theme["ink"], bold=cell.column in centered_columns and cell.value not in (None, ""))
                    cell.fill = PatternFill("solid", fgColor=stripe)
                    horizontal = "center" if cell.column in centered_columns else "left"
                    cell.alignment = Alignment(horizontal=horizontal, vertical="top", wrap_text=True, indent=0 if horizontal == "center" else 1)
                    cell.border = edges(stripe, hairline, rule if theme["style"] == "accent" and cell.column == 1 else None)
                width = max(8, int(widths[cell.column - 1]) - 3)
                lines = max(lines, sum(max(1, len(textwrap.wrap(part, width)))
                                       for part in str(cell.value or "").split("\n")))
            ws.row_dimensions[row[0].row].height = min(409, 36 if row[0].row == 1 else max(42 if theme["style"] == "open" else 36, 16 * lines + 14))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.print_title_rows = "1:1"
        ws.print_area = ws.dimensions
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_setup.horizontalCentered = True
        ws.page_margins = PageMargins(left=0.45, right=0.45, top=0.7, bottom=0.55, header=0.28, footer=0.28)
        ws.oddHeader.left.text = programme_title[:90]
        ws.oddHeader.left.font = "Calibri"
        ws.oddHeader.left.size = 10
        ws.oddFooter.left.text = "Client curriculum"
        ws.oddFooter.left.font = "Calibri"
        ws.oddFooter.left.size = 9
        ws.oddFooter.right.text = "Page &P of &N"
        ws.oddFooter.right.font = "Calibri"
        ws.oddFooter.right.size = 9
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 110
        return ws

    original_days = toc.get("days") or []
    days = []
    for day in original_days:
        if day.get("modules"):
            for module in day["modules"]:
                days.append({**module, "day": day.get("day"), "date": day.get("date"),
                    "timing": module.get("timing") or day.get("timing"), "focus_area": module.get("title") or module.get("topic"),
                    "hours_per_day": None})
        else:
            days.append(day)
    def hours(day):
        minutes = day.get("minutes")
        try:
            value = float(minutes) / 60 if minutes is not None else day.get("hours_per_day", toc.get("hours_per_day"))
            value = float(value)
            return value if math.isfinite(value) and value > 0 else None
        except (TypeError, ValueError):
            return None

    durations = [hours(day) for day in days]
    brief = toc.get("requirement_brief") or {}
    overview = sheet("Program Overview", ["Programme", "Details"], [
        ["Title", ("DRAFT - FOR REVIEW: " if toc.get("draft_export") else "") + str(toc.get("title") or toc.get("domain") or "Training Programme")],
        ["Overview", text(toc.get("overview"))],
        ["Audience", brief.get("audience") or toc.get("audience_level") or toc.get("level") or "To be confirmed"],
        ["Training days", len(original_days)],
        ["Planned training hours", sum(durations) if all(v is not None for v in durations) else "To be confirmed"],
        ["Delivery mode", toc.get("mode") or "To be confirmed"],
        ["Prerequisites", text(toc.get("prerequisites")) or "To be confirmed"],
        ["Learning outcomes", text(toc.get("learning_outcomes"))],
    ], [28, 105])

    plan_rows, outcomes, assessments, delivery_rows = [], [], [], []
    for index, day in enumerate(days, 1):
        title = day.get("focus_area") or day.get("module") or day.get("title") or f"Module {index}"
        topics = text(day.get("subtopics") or (day.get("morning_session") or {}).get("topics"))
        lab = text(day.get("lab") or day.get("lab_task"))
        objectives = text(day.get("learning_objectives"))
        assessment = text(day.get("assessment")) or "To be confirmed"
        deliverable = text(day.get("deliverable") or day.get("lab_deliverable")) or "To be confirmed"
        duration = durations[index - 1]
        row = [day.get("day") or index, day.get("date") or "To be confirmed",
               day.get("timing") or toc.get("timing") or "To be confirmed", title,
               duration, topics, lab, objectives, deliverable, assessment, text(day.get("tools"))]
        plan_rows.append(row)
        outcomes.append([row[0], title, objectives])
        assessments.append([row[0], title, assessment, deliverable])
        steps = day.get("delivery_steps") or []
        if steps:
            for sequence, step in enumerate(steps, 1):
                delivery_rows.append([row[0], title, sequence, text(step.get("method")),
                    step.get("minutes"), text(step.get("activity")), text(step.get("evidence"))])
        else:
            delivery_rows.append([row[0], title, None, "To be confirmed", None,
                "Teaching sequence and activity timings require trainer confirmation.", "To be confirmed"])

    day_headers = ["Day", "Date", "Session timing", "Module", "Hours", "Topics", "Hands-on lab", "Learning outcomes", "Deliverable", "Success measure"]
    technical_headers = ["Day", "Date", "Session timing", "Module", "Hours", "Technical coverage", "Practical lab", "Learning outcomes", "Tools"]
    skills_headers = ["Day", "Date", "Session timing", "Module", "Hours", "Focus areas", "Practical application", "Participant outcomes", "Participant output"]
    syllabus_headers = ["Day", "Date", "Session timing", "Module", "Hours", "Detailed topics", "Lab", "Learning objectives", "Lab deliverable", "Assessment", "Tools"]
    schedules = {
        "execution_plan": ("Day-wise Plan", day_headers, [row[:10] for row in plan_rows], [8, 17, 24, 36, 10, 48, 50, 48, 36, 40]),
        "technical_plan": ("Day-wise Training Plan", technical_headers, [row[:8] + [row[10]] for row in plan_rows], [8, 17, 24, 36, 10, 58, 52, 48, 26]),
        "skills_matrix": ("Skills Matrix", skills_headers, [row[:9] for row in plan_rows], [8, 17, 24, 38, 10, 55, 50, 48, 38]),
        "detailed_syllabus": ("Detailed Syllabus", syllabus_headers, plan_rows, [8, 17, 24, 38, 10, 58, 52, 48, 36, 40, 26]),
        "stone_plan": ("Day Plan", day_headers, [row[:10] for row in plan_rows], [8, 17, 24, 36, 10, 48, 50, 48, 36, 40]),
        "moss_plan": ("Practice Plan", technical_headers, [row[:8] + [row[10]] for row in plan_rows], [8, 17, 24, 36, 10, 58, 52, 48, 26]),
        "sand_plan": ("Skills Outline", skills_headers, [row[:9] for row in plan_rows], [8, 17, 24, 38, 10, 55, 50, 48, 38]),
        "bark_plan": ("Workshop Syllabus", syllabus_headers, plan_rows, [8, 17, 24, 38, 10, 58, 52, 48, 36, 40, 26]),
        "mist_plan": ("Session Plan", day_headers, [row[:10] for row in plan_rows], [8, 17, 24, 36, 10, 48, 50, 48, 36, 40]),
        "olive_plan": ("Learning Syllabus", syllabus_headers, plan_rows, [8, 17, 24, 38, 10, 58, 52, 48, 36, 40, 26]),
    }
    schedule_name, schedule_headers, schedule_rows, schedule_widths = schedules[layout]
    schedule = sheet(schedule_name, schedule_headers, schedule_rows, schedule_widths)
    for row in range(2, schedule.max_row + 1):
        schedule.cell(row, 5).number_format = "0.##"
    if days and all(value is not None for value in durations):
        overview.cell(6, 2, f"=SUM('{schedule.title}'!E2:E{schedule.max_row})")
        overview.cell(6, 2).number_format = "0.##"
    sheet("Module Outcomes", ["Day", "Module", "Learning outcomes"], outcomes, [8, 42, 100])
    sheet("Assessment", ["Day", "Module", "Assessment / success measure", "Deliverable"], assessments, [8, 40, 85, 65])
    sheet("Training Delivery", ["Day", "Module", "Step", "Teaching method", "Minutes", "Trainer / participant activity", "Evidence of completion"],
          delivery_rows, [8, 38, 8, 25, 12, 75, 65])
    check_rows = []
    for row, day in zip(plan_rows, days):
        for check in day.get("acceptance_checks") or []:
            if isinstance(check, dict):
                check_rows.append([row[0], row[3], text(check.get("input_or_condition")),
                    text(check.get("expected_result")), text(check.get("evidence"))])
            elif isinstance(check, str) and check.strip():
                check_rows.append([row[0], row[3], check, "See acceptance statement", "Trainer to confirm evidence"])
    if check_rows:
        sheet("Acceptance Checks", ["Day", "Module", "Input / condition", "Expected result", "Evidence"],
              check_rows, [8, 38, 65, 65, 55])
    outcome_rows = []
    for outcome in toc.get("outcome_coverage") or []:
        for module in outcome.get("modules") or []:
            outcome_rows.append([text(outcome.get("client_outcome")), module.get("day"),
                                 text(module.get("module")), text(module.get("deliverable"))])
    if outcome_rows:
        sheet("Client Objectives", ["Client outcome", "Day", "Module", "Participant deliverable"],
              outcome_rows, [65, 8, 45, 65])
    scenario_rows = []
    for row, day in zip(plan_rows, days):
        if day.get("scenarios"):
            for scenario in day["scenarios"]:
                scenario_rows.append([row[0], row[3], text(scenario.get("situation")),
                    text(scenario.get("activity")), text(scenario.get("outcome")), text(scenario.get("evidence"))])
        elif day.get("scenario"):
            scenario_rows.append([row[0], row[3], text(day["scenario"]), row[6], row[7], row[9]])
    if scenario_rows:
        sheet("Scenarios", ["Day", "Module", "Scenario", "Participant activity", "Expected outcome", "Evidence"],
              scenario_rows, [8, 38, 60, 60, 55, 45])
    readiness = []
    if toc.get("draft_export"):
        quality = toc.get("quality") or {}
        readiness.append(["Draft status", "For review only; this export does not approve client delivery."])
        readiness.extend(["Validation issue", text(v)] for v in quality.get("validation_errors") or [])
        readiness.extend(["Review item", text(v)] for v in quality.get("review_warnings") or [])
        for day in days:
            if day.get("proposed_fields"):
                readiness.append(["Proposed module content", text(day.get("title") or day.get("focus_area")) + ": " + ", ".join(day["proposed_fields"])])
    for item in toc.get("prerequisites") or []:
        readiness.append(["Prerequisite", text(item)])
    for item in toc.get("assumptions") or []:
        readiness.append(["Assumption to confirm", text(item)])
    for item in toc.get("clarification_questions") or []:
        readiness.append(["Clarification", text(item)])
    for item in brief.get("constraints") or []:
        readiness.append(["Client constraint", text(item)])
    for day in days:
        title = day.get("title") or day.get("focus_area") or "Module"
        for item in day.get("preparation_requirements") or []:
            readiness.append(["Trainer preparation: " + str(title), text(item)])
        for item in day.get("prerequisites") or []:
            readiness.append(["Module prerequisite: " + str(title), text(item)])
    if readiness:
        sheet("Readiness & Risks", ["Category", "Requirement / action"], readiness, [30, 110])
    workbook.properties.title = programme_title[:120]
    workbook.properties.subject = "Training curriculum"
    workbook.properties.category = layout.replace("_", " ").title()
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
