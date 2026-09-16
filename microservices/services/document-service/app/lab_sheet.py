"""Consolidate verified cloud estimates without changing their calculations."""
import io
import re
from copy import copy

import openpyxl
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.properties import CalcProperties


def _relocate_formula(formula, source_sheet, offsets):
    tokens = Tokenizer(formula).items
    for token in tokens:
        if token.type != 'OPERAND' or token.subtype != 'RANGE':
            continue
        reference = token.value
        sheet = source_sheet
        if '!' in reference:
            sheet, reference = reference.rsplit('!', 1)
            sheet = sheet.strip("'").replace("''", "'")
        if sheet not in offsets:
            raise ValueError('Unknown worksheet in lab formula: ' + sheet)
        endpoints = reference.split(':')
        relocated = []
        for endpoint in endpoints:
            match = re.fullmatch(r'(\$?[A-Z]+)(\$?)([1-9][0-9]*)', endpoint)
            if not match:
                raise ValueError('Unsupported lab formula reference: ' + token.value)
            column, absolute, row = match.groups()
            relocated.append(f'{column}{absolute}{int(row) + offsets[sheet]}')
        token.value = ':'.join(relocated)
    return '=' + ''.join(token.value for token in tokens)


def combine_lab_estimates(estimates):
    """One worksheet: cloud totals and line items first, expandable evidence below.

    Each input retains its own assumptions and rate card. AWS and Azure are
    shown separately, never silently added as though both clouds are required.
    """
    if not estimates or len(estimates) > 3:
        raise ValueError('Supply between one and three cloud estimates')
    loaded = []
    seen = set()
    for provider, content in estimates:
        if provider not in {'aws', 'azure', 'gcp'} or provider in seen:
            raise ValueError('Invalid or duplicate cloud provider')
        seen.add(provider)
        book = openpyxl.load_workbook(io.BytesIO(content))
        if not {'Client Estimate', 'Resource Cost Breakdown', 'Assumptions', 'Rate Card'} <= set(book.sheetnames):
            raise ValueError('Incomplete lab estimate')
        loaded.append((provider, book))
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'Lab Cost Estimate'
    sheet.append(['Cloud lab cost estimate'])
    sheet.append(['Provider totals are separate. Use the cloud(s) required by the course; do not add alternatives.'])
    sheet.merge_cells('A1:K1')
    sheet.merge_cells('A2:K2')
    blocks = []
    offsets = {provider: {} for provider, _ in loaded}
    row = 5
    # Keep only the quote and resource costs expanded. All supporting inputs
    # remain editable on this same worksheet, with no cross-sheet references.
    # Present the client quote before the detailed technical evidence.
    preferred_order = ('Client Estimate', 'Resource Cost Breakdown', 'TOC Mapping',
                       'Assumptions', 'Rate Card')
    for detail in (False, True):
        for provider, book in loaded:
            source_order = tuple(name for name in preferred_order if name in book.sheetnames) + tuple(
                name for name in book.sheetnames if name not in preferred_order
            )
            for source_name in source_order:
                source = book[source_name]
                is_detail = source.title not in {'Client Estimate', 'Resource Cost Breakdown'}
                if is_detail != detail:
                    continue
                offsets[provider][source.title] = row
                blocks.append((provider, source, row, detail))
                row += source.max_row + 3
    for provider, source, offset, detail in blocks:
        label = sheet.cell(offset, 1, f'{provider.upper()} — {source.title}')
        label.font = Font(bold=True, color='FFFFFF', size=12)
        label.fill = PatternFill('solid', fgColor='17365D')
        sheet.merge_cells(start_row=offset, start_column=1, end_row=offset, end_column=11)
        for source_row in source:
            for cell in source_row:
                if cell.value is None:
                    continue
                value = cell.value
                if cell.data_type == 'f':
                    value = _relocate_formula(value, source.title, offsets[provider])
                target = sheet.cell(cell.row + offset, cell.column, value)
                target.font = copy(cell.font)
                target.fill = copy(cell.fill)
                target.border = copy(cell.border)
                target.number_format = cell.number_format
                target.protection = copy(cell.protection)
                target.alignment = Alignment(vertical='top', wrap_text=True)
                if cell.comment:
                    target.comment = copy(cell.comment)
                if cell.hyperlink:
                    target.hyperlink = copy(cell.hyperlink)
        for merged in source.merged_cells.ranges:
            sheet.merge_cells(start_row=merged.min_row + offset, start_column=merged.min_col,
                              end_row=merged.max_row + offset, end_column=merged.max_col)
        for number in range(1, source.max_row + 1):
            source_height = source.row_dimensions[number].height or 20
            sheet.row_dimensions[number + offset].height = max(source_height, 20)
        if detail:
            sheet.row_dimensions.group(offset + 1, offset + source.max_row, hidden=True)
    widths = {'A': 34, 'B': 18, 'C': 16, 'D': 16, 'E': 18,
              'F': 14, 'G': 38, 'H': 18, 'I': 14, 'J': 18, 'K': 12}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[1].height = 30
    sheet.row_dimensions[2].height = 28
    sheet.freeze_panes = 'A5'
    sheet['A1'].font = Font(bold=True, color='FFFFFF', size=16)
    sheet['A1'].fill = PatternFill('solid', fgColor='0F766E')
    sheet['A1'].alignment = Alignment(vertical='center')
    sheet['A2'].font = Font(italic=True, color='475569', size=10)
    sheet['A2'].alignment = Alignment(vertical='center', wrap_text=True)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.outlinePr.summaryBelow = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_area = sheet.dimensions
    workbook.calculation = CalcProperties(calcMode='auto', fullCalcOnLoad=True, forceFullCalc=True)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
