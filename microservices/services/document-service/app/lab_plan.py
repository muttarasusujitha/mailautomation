"""A concise client view over the existing, auditable lab-cost calculations."""
import json

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _text(value):
    if isinstance(value, (list, tuple)):
        return "\n".join(filter(None, (_text(item) for item in value)))
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    return str(value or "").strip()


def _excel_literal(value):
    return '"' + str(value).replace('"', '""') + '"'


def _vm_description(values, profile):
    """Describe the selected provider product; never guess hardware from a tier."""
    rate = (values.get('rate_card_overrides') or {}).get('VM ' + profile) or {}
    spec = rate.get('specifications') or {}
    parts = [profile + ' VM']
    if spec.get('instance_type'):
        parts.append(str(spec['instance_type']))
    elif rate.get('sku') and rate['sku'] != 'NOT_BILLED':
        parts.append('SKU ' + str(rate['sku']))
    else:
        parts.append('SKU unconfirmed')
    parts.append(str(spec['vcpu']) + ' vCPU' if spec.get('vcpu') else 'vCPU unconfirmed')
    parts.append(str(spec['memory']) + ' RAM' if spec.get('memory') else 'RAM unconfirmed')
    if spec.get('os'):
        parts.append(str(spec['os']))
    return ' / '.join(parts)


def _day_setup(day, index, values):
    setups = values.get('lab_day_setups') or {}
    override = setups[index - 1] if isinstance(setups, list) and index <= len(setups) else (
        setups.get(str(index), setups.get(index)) if isinstance(setups, dict) else None)
    explicit = override or day.get('lab_setup') or values.get('lab_setup')
    if explicit:
        return str(explicit).strip().title()
    from shared.lab_curriculum import lab_setup_kind
    setup = lab_setup_kind(day)
    if setup == 'local':
        return 'Local (from TOC)'
    if setup == 'mixed':
        return 'Mixed local / cloud'
    return None


def _tools_for_day(day):
    """Show only tools explicitly named by the TOC."""
    found = []

    def add(value):
        if isinstance(value, (list, tuple)):
            for item in value:
                add(item)
            return
        value = _text(value)
        if value and value.lower() not in {item.lower() for item in found}:
            found.append(value)

    for key in ('tools', 'tool', 'platforms', 'technologies', 'technology'):
        add(day.get(key))
    for module in day.get('modules') or []:
        if isinstance(module, dict):
            for key in ('tools', 'tool', 'platforms', 'technologies', 'technology'):
                add(module.get(key))
    scope = _text({key: day.get(key) for key in ('title', 'topic', 'subtopics', 'lab', 'lab_task')}).lower()
    for name in ('AWS IAM', 'Amazon VPC', 'Amazon EC2', 'Amazon S3', 'CloudTrail',
                 'Azure Portal', 'Azure VM', 'Azure Storage', 'Azure Monitor',
                 'Amazon EKS', 'Amazon RDS', 'Elastic Load Balancing', 'AWS Lambda',
                 'Amazon DynamoDB', 'Amazon ECR', 'AWS Secrets Manager', 'CloudTrail',
                 'Azure Kubernetes Service (AKS)',
                 'Azure SQL', 'Azure Blob Storage', 'Azure Load Balancer',
                 'Azure Application Gateway', 'Azure Functions', 'Azure Cosmos DB',
                 'Azure App Service', 'Azure Key Vault',
                 'Docker', 'Kubernetes', 'Terraform', 'Ansible', 'Jenkins',
                 'Git', 'Linux shell', 'Python', 'SQL'):
        if name.lower() in scope:
            add(name)
    return ', '.join(found) or 'See TOC activity'


def add_lab_plan(workbook, days, values, profile_rate_rows):
    """Link daily infrastructure and the compact summary to the existing model.

    Daily costs exclude course-wide services and commercial charges, which are
    shown separately so no arbitrary allocation obscures local training days.
    This presentation does not certify caller-provided prices or alter billing.
    """
    sheet = workbook.create_sheet("Lab Plan", 0)
    verified = (
        values.get("pricing_status") == "provider_api_verified_public_retail"
        and bool(values.get("rate_snapshot_id"))
        and bool(values.get("rate_checked_at"))
        and bool(values.get("fx_rate_source"))
    )
    money_format = '"INR "#,##0.00;[Red]("INR "#,##0.00);"INR "0.00'
    unresolved = [label for key, label in (
        ('local_costs_status', 'local hardware/electricity'),
        ('license_costs_status', 'required licenses'))
        if values.get(key) not in ('covered', 'not_required')]
    unresolved.extend(values.get('required_unpriced_costs') or [])
    incomplete = not verified or bool(unresolved)

    def put(row, column, value):
        cell = sheet.cell(row, column, value)
        # Course titles and practical activities are untrusted text, not Excel formulas.
        if isinstance(value, str) and not value.startswith("="):
            cell.data_type = "s"
        return cell

    def text(row, column, value):
        cell = put(row, column, value)
        cell.data_type = "s"
        return cell

    def band(row, value, color="17365D"):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        cell = text(row, 1, value)
        cell.fill = PatternFill("solid", fgColor=color)
        cell.font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
        sheet.row_dimensions[row].height = 26

    def note(row, label, value, height=34, formula=False):
        text(row, 1, label)
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)
        (put if formula else text)(row, 2, value)
        sheet.row_dimensions[row].height = height

    def amount(row, label, formula):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        text(row, 1, label)
        cell = put(row, 7, formula if verified else "Unverified")
        cell.number_format = money_format

    band(1, "Training lab cost estimate", "0F766E")
    note(2, "Training", "='Assumptions'!B4", formula=True)
    note(3, "Delivery", "='Assumptions'!B5&\" | Region: \"&'Assumptions'!B6&\" | \"&'Assumptions'!B9&\" participants | \"&'Assumptions'!B7&\" hours/day | \"&'Assumptions'!B8&\" days\"", formula=True)
    status = 'Verified cloud rates and FX' if verified else 'Unverified cloud rates or FX'
    if incomplete:
        status = 'Incomplete estimate - ' + status
        if unresolved:
            status += '; unresolved: ' + ', '.join(unresolved)
    else:
        status += '; local costs/licenses confirmed covered or not required. Review usage assumptions.'
    note(4, "Live price status", status, 38)
    sheet.merge_cells(start_row=5, start_column=1, end_row=5, end_column=5)
    text(5, 1, "ESTIMATED LAB COST — PLANNED USAGE" if not incomplete else "PRICED SUBTOTAL — INCOMPLETE ESTIMATE")
    total_cell = put(5, 6, "='Client Estimate'!B12" if verified else "Unverified")
    sheet.merge_cells(start_row=5, start_column=6, end_row=5, end_column=7)
    total_cell.number_format = money_format
    total_cell.font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    total_cell.fill = PatternFill("solid", fgColor="0F766E")
    sheet["A5"].font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    sheet["A5"].fill = PatternFill("solid", fgColor="0F766E")
    sheet.row_dimensions[5].height = 30
    headers = ("Day", "TOC topic", "Practical activity", "Tools used", "Resources / count", "Hours", "Estimated cost (INR)")
    for column, label in enumerate(headers, 1):
        text(6, column, label)
    sheet.row_dimensions[6].height = 32

    for index, day in enumerate(days, 1):
        row, mapping_row = index + 6, index + 3
        ref = lambda column: f"'TOC Mapping'!{column}{mapping_row}"
        text(row, 1, _text(day.get("day") or index))
        no_cloud = f"SUM({ref('C')},{ref('E')}:H{mapping_row})=0"
        setup = _day_setup(day, index, values)
        if setup:
            setup_formula = _excel_literal(setup)
            if setup.lower().startswith('local'):
                setup_formula = f'IF({no_cloud},{setup_formula},"Mixed - local setup with cloud resources; review")'
        else:
            # Counts indicate a proposed allocation, not confirmed deployment.
            setup_formula = (f'IF({no_cloud},"Setup unconfirmed - no cloud resources",'
                f'IF({ref("C")}=\'Assumptions\'!$B$9,IF(SUM({ref("E")}:H{mapping_row})>0,'
                '"Individual + shared (assumed)","Individual (assumed)"),"Shared (assumed)"))')
        put(row, 2, f'={ref("B")}&CHAR(10)&"Setup: "&{setup_formula}')
        practical = next((day.get(key) for key in ("lab", "lab_task", "practical_lab", "hands_on") if day.get(key)), None)
        if day.get('modules'):
            practical = '\n'.join(
                str(module.get('title') or module.get('topic') or 'Module') + ': ' +
                str(module.get('lab') or module.get('lab_task') or 'Lab details to be confirmed')
                for module in day['modules'])
        text(row, 3, _text(practical) or "Practical activity to be confirmed from TOC")
        # Each populated resource has a matching quantity line, with units.
        resources = [
            ("C", f'IF({ref("D")}="Heavy",{_excel_literal(_vm_description(values, "Heavy"))},{_excel_literal(_vm_description(values, "Light"))})', f'{ref("C")}&" VMs"'),
            ("E", '"Kubernetes control plane"', f'{ref("E")}&" clusters"'),
            ("F", '"Kubernetes workers"', f'{ref("F")}&" nodes"'),
            ("G", '"Managed database"', f'{ref("G")}&" instances"'),
            ("H", '"Object storage"', f'{ref("H")}&" GB"'),
        ]
        resource_parts, quantity_parts = [], []
        for column, label, quantity in resources:
            resource_parts.append(f'IF({ref(column)}>0,{label}&CHAR(10),"")')
            quantity_parts.append(f'IF({ref(column)}>0,{quantity}&CHAR(10),"")')
        nodes = f"({ref('C')}+{ref('F')})"
        resource_parts.append(f'IF({nodes}>0,"Attached disks","")')
        quantity_parts.append(f'IF({nodes}>0,{nodes}*\'Assumptions\'!$B$20&" GB","")')
        empty = f"SUM({ref('C')},{ref('E')}:H{mapping_row})=0"
        text(row, 4, _tools_for_day(day))
        put(row, 5, f'=IF({empty},"No mapped cloud resources",' + '&'.join(resource_parts) + ')&CHAR(10)&IF(' + empty + ',"","Count: "&' + '&'.join(quantity_parts) + ')')
        put(row, 6, f"={ref('I')}*'Assumptions'!$B$7")
        vm_rate = f'IF({ref("D")}="Heavy",\'Rate Card\'!$F${profile_rate_rows["Heavy"]},\'Rate Card\'!$F${profile_rate_rows["Light"]})'
        hourly = f"({ref('C')}*{vm_rate}+{ref('E')}*'Resource Cost Breakdown'!$D$5+{ref('F')}*'Resource Cost Breakdown'!$D$6+{ref('G')}*'Resource Cost Breakdown'!$D$11)*F{row}"
        retained = f"({nodes}*'Assumptions'!$B$20*'Resource Cost Breakdown'!$D$7+{ref('H')}*'Resource Cost Breakdown'!$D$8)*{ref('I')}/30"
        put(row, 7, f"={hourly}+{retained}" if verified else "Unverified")
        sheet.cell(row, 7).number_format = money_format
        # Generous height for resource lines and practical descriptions.
        lines = max(10, sum(max(1, (len(line) + 39) // 40) for line in sheet.cell(row, 3).value.splitlines()))
        sheet.row_dimensions[row].height = min(360, max(150, lines * 15))

    end = 6 + len(days)
    amount(end + 1, "Day-wise cloud infrastructure subtotal", f"=SUM(G7:G{end})")
    amount(end + 2, "Shared course usage: egress, build runner, monitoring",
           "=SUM('Resource Cost Breakdown'!H9:H10)+'Resource Cost Breakdown'!H12")
    band(end + 4, "Price summary")
    summary = [
        ("Cloud infrastructure cost", "='Client Estimate'!B8"),
        ("Support", "='Client Estimate'!B9"),
        ("Contingency", "='Client Estimate'!B10"),
        ("Tax / GST", "='Client Estimate'!B11"),
        ("Service adjustment", "='Client Estimate'!B12-SUM('Client Estimate'!B8:B11)"),
        ("Priced lab subtotal - incomplete estimate" if incomplete else "Total estimated lab cost (confirmed scope)", "='Client Estimate'!B12"),
        ("Priced subtotal per participant" if incomplete else "Cost per participant", "='Client Estimate'!B13"),
    ]
    for offset, (label, formula) in enumerate(summary, end + 5):
        amount(offset, label, formula)
    note(end + 12, "Local / licenses", 'Unverified / excluded: ' + ', '.join(unresolved) + '. All-in total is incomplete.' if unresolved else
         'Local costs: ' + values['local_costs_status'] + '; licenses: ' + values['license_costs_status'] + '. Confirmations apply to this training scope.')
    band(end + 14, "What this estimate includes")
    resource_setup = "\n".join(filter(None, (
        "Overall lab setup: " + _text(values.get("lab_setup")) if values.get("lab_setup") else "",
        _text(values.get("architecture_note")),
    ))) or "Quantities are total provisioned resources from TOC Mapping; confirm individual versus shared setup."
    resource_setup_height = min(360, max(60, 15 * sum(
        max(1, (len(line) + 72) // 73) for line in resource_setup.splitlines())))
    note(end + 15, "Resource setup", resource_setup, resource_setup_height)
    note(end + 16, "Billing", "Hours are per-resource lab runtime for each TOC entry. Disks and object storage accrue for active days, prorated over 30 days. Extra idle runtime and retention beyond active days are excluded; confirm shutdown/deletion.", 48)
    note(end + 17, "Cost basis", "Daily costs show cloud infrastructure only. Course-wide egress, build and monitoring are charged once in the separate shared-usage row. Support, contingency, tax and the existing service adjustment appear in the summary.", 48)
    note(end + 18, "Applied assumptions", "=\"Contingency: \"&TEXT('Assumptions'!B16,\"0.0%\")&\"; tax: \"&TEXT('Assumptions'!B17,\"0.0%\")&\". Review configured/default percentages and support before approval.\"", formula=True)
    note(end + 19, "Excluded costs", "Unmapped services, local hardware, electricity and unpriced licenses are excluded. A complete all-in total requires these costs or confirmation that they are already covered.", 44)
    note(end + 20, "Provider sources", _text(values.get("rate_snapshot_source")) or "Unverified - no quote-specific provider source")
    note(end + 21, "Verified at (UTC)", _text(values.get("rate_checked_at")) or "Unverified")
    note(end + 22, "FX source / date", " | ".join(filter(None, (_text(values.get("fx_rate_source")), _text(values.get("fx_rate_date"))))) or "Unverified")
    note(end + 23, "Price basis", "Estimated lab cost for the planned training usage. Verified provider rates are converted from USD to INR using the fetched exchange rate. The estimate includes applicable support, contingency, tax and the configured service increase. Actual AWS/Azure charges may differ if runtime, storage or other usage changes.", 60)

    for cells in sheet:
        for cell in cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row not in (1, end + 4, end + 14):
                cell.font = Font(name="Calibri", size=11, color="1E293B", bold=cell.row == 6)
            if cell.row == 6:
                cell.fill = PatternFill("solid", fgColor="DCEFEA")
            elif 7 <= cell.row <= end and cell.row % 2:
                cell.fill = PatternFill("solid", fgColor="F1F5F9")
    for column, width in enumerate((10, 28, 42, 25, 34, 12, 25), 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "C7"
    sheet.auto_filter.ref = f"A6:G{end}"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:6"
    sheet.print_area = f"A1:G{end + 23}"
    return sheet
