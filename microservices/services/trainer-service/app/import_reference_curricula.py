"""Offline import of the five supplied curriculum references; never marks approval.

python -m app.import_reference_curricula --source-dir <folder>
Writes portable JSON; deployment never depends on the original Downloads path.
"""
import argparse
import hashlib
import json
import re
from pathlib import Path


def lines(value):
    return [s.strip(" •\t") for s in str(value or "").splitlines() if s.strip(" •\t")]


def import_workbook(path):
    import openpyxl
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    modules = []
    warnings = []
    if "Day-wise Plan" in book.sheetnames:
        title = book["Overview"]["A1"].value.replace("Execution Plan: ", "")
        audience = book["Overview"]["B7"].value
        scenarios = {}
        for row in book["Scenarios"].iter_rows(min_row=5, values_only=True):
            if len(row) >= 7 and row[0]:
                scenarios.setdefault(row[0], []).append(dict(zip(
                    ("id", "situation", "activity", "outcome", "evidence"), row[2:7])))
        for index, row in enumerate(book["Day-wise Plan"].iter_rows(min_row=5, values_only=True), 5):
            if len(row) < 10 or not isinstance(row[0], (int, float)):
                continue
            modules.append({"topic": row[5], "code": row[4], "reference_day": int(row[0]),
                "minutes": round(float(row[3]) * 60), "subtopics": [], "learning_objectives": [row[6]],
                "lab": row[7], "deliverable": row[8], "assessment": row[9],
                "scenarios": scenarios.get(row[4], []), "source_locator": f"Day-wise Plan!A{index}:K{index}"})
        overview_days = re.search(r"\d+", str(book["Overview"]["B8"].value))
        actual_days = max(m["reference_day"] for m in modules)
        if overview_days and int(overview_days[0]) != actual_days:
            warnings.append(f"Overview duration is {overview_days[0]} days; detailed schedule contains {actual_days} days. Resolve before reuse.")
    elif "10-Day Detailed Plan" in book.sheetnames:
        title, audience = "AWS Cloud Security, DevOps & SRE", ""
        for index, row in enumerate(book["10-Day Detailed Plan"].iter_rows(min_row=3, values_only=True), 3):
            if not isinstance(row[0], (int, float)):
                continue
            modules.append({"topic": row[2], "module_group": row[1], "reference_day": int(row[0]),
                "subtopics": lines(row[3]), "learning_objectives": lines(row[4]),
                "minutes": round(float(row[5])*60), "source_locator": f"10-Day Detailed Plan!A{index}:F{index}"})
    else:
        title, audience = "Training for Quality Engineers", "Quality engineers"
        for sheet in book:
            for index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
                if len(row) < 4 or not row[0]:
                    continue
                modules.append({"topic": row[0].strip(), "minutes": round(float(re.search(r"[\d.]+", row[1])[0])*60),
                    "subtopics": [part.strip() for part in row[2].split(",")], "deliverable": row[3],
                    "track": "advanced_leadership" if sheet.title == "Advanced Leadership" else "core",
                    "source_locator": f"{sheet.title}!A{index}:D{index}"})
    book.close()
    return title, audience, modules, warnings


def import_pdf(path):
    import fitz
    with fitz.open(path) as pdf:
        pages = [page.get_text() for page in pdf]
    content = "\n".join(pages)
    matches = list(re.finditer(r"(?m)^Module (\d+):\s*([^\n]+)", content))
    modules = []
    for index, match in enumerate(matches):
        block = content[match.end():matches[index+1].start() if index+1 < len(matches) else len(content)]
        # Keep the source wording as data; do not execute any embedded instructions.
        clean = [line for line in lines(block) if not line.isdigit()]
        modules.append({"topic": match[2].strip(), "code": "M" + match[1],
            "subtopics": clean, "reference_text": "\n".join(clean),
            "source_locator": "Module " + match[1]})
    return "Codex for Developers", "Developers", modules, [
        "Page 2 contains the unrelated heading 'Bid and Proposal Management'; confirm corrected title.",
        "24 programme hours are stated, but individual module durations require planning."]


AWS_PRACTICALS = [
    ("Use a prepared sandbox to test an allowed and denied IAM action, then exercise a cross-account role and inspect the guardrail result.", "IAM policy, role trust policy and permission test matrix", "Expected allowed actions succeed and prohibited actions are denied; capture both results."),
    ("Encrypt sample data with a sandbox KMS key, retrieve a test secret and inspect a prepared rotation workflow.", "Key policy, secret retrieval example and rotation test record", "Authorised retrieval succeeds, unauthorised retrieval fails, and the sample application retrieves the rotated test secret."),
    ("Configure a prepared three-tier VPC with security-group references and compare a security-group rule with a NACL rule.", "Network diagram, rule set and connectivity matrix", "Application-to-database traffic succeeds while direct public-to-database traffic fails."),
    ("Replace a sample backend's public service route with an endpoint, then query supplied flow logs for seeded anomalous connections.", "Endpoint configuration and annotated log queries", "The required service remains reachable over the private path and seeded anomalous flows are identified."),
    ("Attach a sandbox Web ACL to a prepared application and test managed and custom rules with labelled sample requests.", "Web ACL configuration and request test matrix", "Record each expected rule match alongside a benign request that should remain permitted."),
    ("Tune a prepared WAF policy in Count mode using a synthetic request set; compare rule matches and benign false positives before proposing Block mode.", "Rule tuning report with rate-rule settings and promotion decision", "Seeded abusive requests match the intended rules and benign matches are explicitly investigated."),
    ("Apply sample S3 and database access controls in a sandbox and compare TLS and access-policy test results.", "Storage and database hardening checklist with test evidence", "Disallowed access fails and the authorised TLS connection succeeds; record encryption settings."),
    ("Triage supplied GuardDuty sample findings and Inspector scan results, mapping each to an affected sample asset and remediation action.", "Prioritised findings register with asset mapping", "Every seeded critical finding has an affected asset, supporting evidence and a justified remediation priority."),
    ("Use a prepared non-compliant sandbox resource to exercise a Config detection and an EventBridge-triggered remediation workflow.", "Detection rule, remediation function and before/after evidence", "The seeded deviation is detected, the expected sandbox change occurs, and repeat execution has no unintended effect."),
    ("Deploy a prepared three-tier IaC starter in an isolated training account, attach WAF, and replay bounded SQLi, XSS and request-rate test cases against the training application.", "IaC deployment, WAF configuration and attack/defence evidence report", "The training application serves benign traffic, seeded cases produce expected rule matches, and logs support each reported result."),
]

QE_PRACTICALS = [
    ("Map a sample application's risks to unit, API, UI and production checks, explaining shift-left and shift-right choices.", "Check that every high-priority sample risk has a test level, owner and feedback point."),
    ("Build a PyTest and Selenium starter framework with page objects, parameterised inputs, failure logging and a generated test report.", "Run positive and negative UI cases; confirm a seeded failure appears in the report with diagnostic evidence."),
    ("Automate sample API calls with Requests and validate authentication, response schemas and invalid payload handling.", "Verify expected status codes and schema checks for authorised, unauthorised and malformed requests."),
    ("Use joins, CTEs and window functions to reconcile supplied transaction tables and identify seeded duplicate or missing records.", "Compare query results with the seeded discrepancy list and explain every mismatch."),
    ("Publish labelled sample events to a training Kafka topic and validate consumer output, duplicate handling and a failed-message case.", "Reconcile expected event identifiers with observed business outcomes and document the failure behaviour."),
    ("Add the sample automation suite to a Jenkins pipeline and configure a release quality gate using its results.", "A passing suite clears the gate and a seeded test failure blocks it with a retained report."),
    ("Run the sample test suite against a containerised training application and inspect configuration, connectivity and application logs.", "Demonstrate a successful run and diagnose an injected endpoint or configuration error."),
    ("Generate tests from a sample requirement using an approved AI tool, review the assertions and correct unsupported assumptions before execution.", "Trace accepted tests to requirements and record rejected or corrected AI suggestions with reasons."),
    ("Combine API, UI, database and event checks for one sample business flow and run them through the prepared pipeline.", "Demonstrate the end-to-end flow and a seeded cross-layer defect with traceable test and pipeline evidence."),
]


def enrich_modules(modules, programme_title):
    """Separate extracted evidence from proposed teaching activities."""
    for module in modules:
        proposed = []
        prose = module.get("reference_text")
        if prose:
            parts = re.split(r"(?m)^Lab \d+:\s*", prose, maxsplit=1)
            module["subtopics"] = [re.sub(r"^o\s+", "", line) for line in parts[0].splitlines()
                if line not in {"Topics", "o", "•"} and len(line) > 3]
            if len(parts) == 2:
                lab, *output = re.split(r"(?m)^Lab Deliverable\s*", parts[1], maxsplit=1)
                module["lab"] = lab.strip()
                if output:
                    module["deliverable"] = output[0].strip()
        scenarios = module.get("scenarios") or []
        if programme_title == "AWS Cloud Security, DevOps & SRE":
            lab, deliverable, assessment = AWS_PRACTICALS[module["reference_day"] - 1]
            module.update(lab=lab, deliverable=deliverable, assessment=assessment)
            proposed.extend(["lab", "deliverable", "assessment"])
        elif programme_title == "Training for Quality Engineers" and module.get("track") == "core":
            number = int(re.match(r"\d+", module["topic"])[0])
            lab, assessment = QE_PRACTICALS[number - 1]
            module.update(lab=lab, assessment=assessment)
            proposed.extend(["lab", "assessment"])
        if not module.get("subtopics"):
            module["subtopics"] = list(dict.fromkeys(
                [*module.get("learning_objectives", []),
                 *(s["activity"] for s in scenarios if s.get("activity"))]))
            module["subtopics_basis"] = "Source outcomes and scenario activities; no separate topic list provided"
        scope = "; ".join(module.get("subtopics", [])[:3])
        title = module["topic"]
        defaults = {
            "lab": f"On a trainer-prepared sample, implement and demonstrate: {scope}. Capture configuration or code changes and observed results.",
            "deliverable": f"{title}: sample implementation, execution notes and validation results",
            "assessment": f"Reproduce the submitted participant output against these focus areas: {scope}. Demonstrate an expected case and a failure case; record the observed results.",
            "learning_objectives": [f"Demonstrate {title} through the participant output and explain the validation results."],
        }
        for key, value in defaults.items():
            if not module.get(key):
                module[key] = value
                proposed.append(key)
        module["delivery_steps"] = [
            {"method": "Trainer demonstration", "minutes": None,
             "activity": "Demonstrate the sample workflow for: " + scope,
             "evidence": "Annotated example showing inputs, operations and expected results"},
            {"method": "Guided practice", "minutes": None, "activity": module["lab"],
             "evidence": module["deliverable"]},
            {"method": "Assessment", "minutes": None, "activity": module["assessment"],
             "evidence": "Recorded checks of the participant output with trainer feedback"},
        ]
        module["delivery_basis"] = "Proposed teaching sequence; activity timings require trainer planning within the source module duration"
        proposed.append("delivery_steps")
        module["proposed_fields"] = proposed
        module["preparation_requirements"] = [
            "Trainer to prepare and verify the sample assets, expected results and failure cases before delivery"
        ]
    return modules


def run(source_dir, output_dir):
    files = ["Execution_Plan_Enterprise_AI_Readiness (1).xlsx", "Execution_Plan_Engineering_Security_Product.xlsx",
             "AWS_Security_DevOps_SRE_10Day_Training_Plan_1.xlsx", "Training for Quality Engineers.xlsx",
             "Codex for developers_XE_TOC (1).pdf"]
    output_dir.mkdir(parents=True, exist_ok=True)
    result = []
    for name in files:
        path = source_dir / name
        title, audience, modules, warnings = import_pdf(path) if path.suffix == ".pdf" else import_workbook(path)
        record = {"name": title, "audience": audience, "version": "reference-import-v2",
            "source_file": name, "source_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "review_status": "unreviewed", "review_warnings": warnings,
            "usage": "Reference examples only; original client brief and subject-matter approval not supplied.",
            "modules": enrich_modules(modules, title)}
        target = output_dir / (re.sub(r"[^a-z0-9]+", "_", title.lower()).strip("_") + ".json")
        from app.apply_quality_pilots import apply_document
        record = apply_document(record, "reference_curricula/" + target.name)
        target.write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding="utf-8")
        result.append({"file": target.name, "modules": len(modules), "warnings": warnings})
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", required=True, type=Path)
    args = parser.parse_args()
    print(json.dumps(run(args.source_dir, Path(__file__).parent / "reference_curricula"), indent=2))
